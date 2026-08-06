from __future__ import annotations

import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from modules.shipments.shipment_repository import ShipmentRepository
from modules.suppliers.supplier_repository import SupplierRepository


class ShipmentFormatDialog(tk.Toplevel):
    """공급처별 회신 송장 엑셀 열 구성을 등록합니다."""

    def __init__(self, parent: tk.Misc) -> None:
        super().__init__(parent)
        self.title("공급처 송장 양식 등록")
        self.geometry("620x360")
        self.resizable(False, False)
        self.transient(parent.winfo_toplevel())
        self.grab_set()

        self.repository = ShipmentRepository()
        self.supplier_repository = SupplierRepository(
            self.repository.database_path
        )
        self.suppliers: dict[str, int] = {}
        self.column_numbers: dict[str, int] = {}
        self.file_path: Path | None = None

        self.supplier_var = tk.StringVar()
        self.file_var = tk.StringVar()
        self.header_row_var = tk.IntVar(value=1)
        self.order_column_var = tk.StringVar()
        self.carrier_column_var = tk.StringVar()
        self.tracking_column_var = tk.StringVar()

        self._build_ui()
        self._load_suppliers()

    def _build_ui(self) -> None:
        frame = ttk.Frame(self, padding=18)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)

        ttk.Label(frame, text="공급처").grid(row=0, column=0, sticky="w")
        self.supplier_combo = ttk.Combobox(
            frame,
            textvariable=self.supplier_var,
            state="readonly",
        )
        self.supplier_combo.grid(
            row=0, column=1, columnspan=2, sticky="ew", pady=(0, 10)
        )
        self.supplier_combo.bind("<<ComboboxSelected>>", self._load_saved)

        ttk.Label(frame, text="Excel 파일").grid(row=1, column=0, sticky="w")
        ttk.Entry(frame, textvariable=self.file_var, state="readonly").grid(
            row=1, column=1, sticky="ew", pady=(0, 10)
        )
        ttk.Button(frame, text="선택", command=self._select_file).grid(
            row=1, column=2, padx=(8, 0), pady=(0, 10)
        )

        ttk.Label(frame, text="헤더 행").grid(row=2, column=0, sticky="w")
        ttk.Spinbox(
            frame,
            from_=1,
            to=9999,
            textvariable=self.header_row_var,
            width=10,
        ).grid(row=2, column=1, sticky="w", pady=(0, 10))
        ttk.Button(frame, text="헤더 불러오기", command=self._load_headers).grid(
            row=2, column=2, padx=(8, 0), pady=(0, 10)
        )

        fields = (
            ("주문번호 열", self.order_column_var),
            ("택배사 열", self.carrier_column_var),
            ("송장번호 열", self.tracking_column_var),
        )
        self.column_combos: list[ttk.Combobox] = []
        for index, (label, variable) in enumerate(fields, start=3):
            ttk.Label(frame, text=label).grid(row=index, column=0, sticky="w")
            combo = ttk.Combobox(frame, textvariable=variable, state="readonly")
            combo.grid(
                row=index,
                column=1,
                columnspan=2,
                sticky="ew",
                pady=(0, 10),
            )
            self.column_combos.append(combo)

        buttons = ttk.Frame(frame)
        buttons.grid(row=6, column=0, columnspan=3, sticky="e", pady=(8, 0))
        ttk.Button(buttons, text="취소", command=self.destroy).pack(
            side="left", padx=(0, 8)
        )
        ttk.Button(buttons, text="저장", command=self._save).pack(side="left")

    def _load_suppliers(self) -> None:
        suppliers = self.supplier_repository.get_suppliers(active_only=True)
        self.suppliers = {
            str(row["supplier_name"]): int(row["id"])
            for row in suppliers
        }
        names = list(self.suppliers)
        self.supplier_combo.configure(values=names)
        if names:
            self.supplier_var.set(names[0])
            self._load_saved()

    def _select_file(self) -> None:
        selected = filedialog.askopenfilename(
            title="공급처 회신 송장 Excel 선택",
            filetypes=[("Excel", "*.xlsx *.xlsm")],
            parent=self,
        )
        if not selected:
            return
        self.file_path = Path(selected)
        self.file_var.set(selected)
        self._load_headers()

    def _load_headers(self) -> None:
        if self.file_path is None:
            messagebox.showwarning("파일 선택", "Excel 파일을 선택하세요.", parent=self)
            return
        try:
            header_row = int(self.header_row_var.get())
            workbook = load_workbook(
                self.file_path,
                data_only=True,
                read_only=True,
            )
            try:
                worksheet = workbook[workbook.sheetnames[0]]
                if header_row < 1 or header_row > worksheet.max_row:
                    raise ValueError("헤더 행이 워크시트 범위를 벗어났습니다.")
                choices: list[str] = []
                self.column_numbers = {}
                for column in range(1, worksheet.max_column + 1):
                    value = worksheet.cell(header_row, column).value
                    header = str(value).strip() if value is not None else ""
                    if not header:
                        continue
                    label = f"{get_column_letter(column)} - {header}"
                    choices.append(label)
                    self.column_numbers[label] = column
                if not choices:
                    raise ValueError("선택한 행에서 헤더를 찾지 못했습니다.")
            finally:
                workbook.close()
        except Exception as exc:
            messagebox.showerror("헤더 오류", str(exc), parent=self)
            return

        for combo in self.column_combos:
            combo.configure(values=choices)

    def _load_saved(self, _event: Any = None) -> None:
        supplier_id = self.suppliers.get(self.supplier_var.get())
        if supplier_id is None:
            return
        mapping = self.repository.get_supplier_shipment_format(supplier_id)
        if mapping:
            self.header_row_var.set(int(mapping["header_row"]))

    def _save(self) -> None:
        supplier_id = self.suppliers.get(self.supplier_var.get())
        if supplier_id is None:
            messagebox.showwarning("공급처", "공급처를 선택하세요.", parent=self)
            return
        if self.file_path is None:
            messagebox.showwarning("파일", "Excel 파일을 선택하세요.", parent=self)
            return

        selected = (
            self.order_column_var.get(),
            self.carrier_column_var.get(),
            self.tracking_column_var.get(),
        )
        columns = [self.column_numbers.get(value) for value in selected]
        if any(column is None for column in columns):
            messagebox.showwarning("열 선택", "필수 열을 모두 선택하세요.", parent=self)
            return
        if len(set(columns)) != 3:
            messagebox.showwarning("열 선택", "각 항목에 서로 다른 열을 선택하세요.", parent=self)
            return

        try:
            header_row = int(self.header_row_var.get())
            workbook = load_workbook(
                self.file_path,
                data_only=True,
                read_only=True,
            )
            try:
                worksheet = workbook[workbook.sheetnames[0]]
                if header_row < 1 or header_row > worksheet.max_row:
                    raise ValueError("헤더 행이 워크시트 범위를 벗어났습니다.")
                for column in columns:
                    if int(column) > worksheet.max_column:
                        raise ValueError("선택한 열이 워크시트 범위를 벗어났습니다.")
                    value = worksheet.cell(header_row, int(column)).value
                    if value is None or not str(value).strip():
                        raise ValueError("선택한 열에 헤더 값이 없습니다.")
            finally:
                workbook.close()

            self.repository.save_supplier_shipment_format(
                supplier_id=supplier_id,
                header_row=header_row,
                order_number_column=int(columns[0]),
                carrier_column=int(columns[1]),
                tracking_number_column=int(columns[2]),
            )
        except Exception as exc:
            messagebox.showerror("저장 오류", str(exc), parent=self)
            return

        messagebox.showinfo("저장 완료", "공급처 송장 양식을 저장했습니다.", parent=self)
        self.destroy()
