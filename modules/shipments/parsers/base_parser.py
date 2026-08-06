from __future__ import annotations

import re
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class BaseShipmentParser(ABC):
    """공급처 또는 표준 송장 엑셀 파서의 공통 기반입니다."""

    parser_key = "base"
    display_name = "기본 파서"
    supplier_name: str | None = None

    @abstractmethod
    def detect(self, file_path: str | Path) -> bool:
        """해당 파일을 이 파서가 처리할 수 있는지 반환합니다."""
        raise NotImplementedError

    @abstractmethod
    def parse(self, file_path: str | Path) -> dict[str, Any]:
        """
        엑셀을 읽고 아래 공통 형태로 반환합니다.

        {
            "source_type": str,
            "supplier_name": str,
            "rows": list[dict],
            "errors": list[dict],
            "total_count": int,
            "valid_count": int,
            "error_count": int,
        }
        """
        raise NotImplementedError

    @staticmethod
    def open_first_sheet(file_path: str | Path):
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"엑셀 파일을 찾을 수 없습니다.\n{path}")

        workbook = load_workbook(
            filename=path,
            data_only=True,
            read_only=True,
        )
        worksheet = workbook[workbook.sheetnames[0]]
        return workbook, worksheet

    @staticmethod
    def clean_text(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value)).strip()
        return str(value).strip()

    @classmethod
    def normalize_phone(cls, value: Any) -> str:
        return re.sub(r"[^0-9]", "", cls.clean_text(value))

    @classmethod
    def normalize_quantity(cls, value: Any) -> int:
        if value is None:
            return 0
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return 0

    @classmethod
    def normalize_carrier(cls, value: Any) -> str:
        original = cls.clean_text(value)
        carrier = original.replace(" ", "")

        aliases = {
            "대한통운": "CJ대한통운",
            "CJ": "CJ대한통운",
            "CJ택배": "CJ대한통운",
            "CJ대한통운": "CJ대한통운",
            "로젠": "로젠택배",
            "로젠택배": "로젠택배",
            "한진": "한진택배",
            "한진택배": "한진택배",
            "롯데": "롯데택배",
            "롯데택배": "롯데택배",
            "롯데글로벌로지스": "롯데택배",
            "우체국": "우체국택배",
            "우체국택배": "우체국택배",
            "경동": "경동택배",
            "경동택배": "경동택배",
        }
        return aliases.get(carrier, original)

    @classmethod
    def clean_tracking_number(cls, value: Any) -> str:
        text = cls.clean_text(value)
        if text.endswith(".0"):
            text = text[:-2]
        return re.sub(r"\s+", "", text)

    @classmethod
    def split_tracking_numbers(cls, value: Any) -> list[str]:
        text = cls.clean_tracking_number(value)
        if not text:
            return []

        numbers: list[str] = []
        for part in re.split(r"[/,\n;]+", text):
            number = re.sub(r"\s+", "", part)
            if number and number not in numbers:
                numbers.append(number)
        return numbers

    @classmethod
    def build_header_map(
        cls,
        worksheet: Any,
        header_row: int = 1,
    ) -> dict[str, int]:
        result: dict[str, int] = {}
        for column_number in range(1, worksheet.max_column + 1):
            header = cls.clean_text(
                worksheet.cell(
                    row=header_row,
                    column=column_number,
                ).value
            )
            if header:
                result[header] = column_number
        return result
