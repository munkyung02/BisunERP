from __future__ import annotations

import shutil
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
DATABASE_PATH = PROJECT_ROOT / "data" / "bisun_erp.db"


class ExportService:
    """ERP 데이터를 엑셀로 내보내고 DB 백업을 생성합니다."""

    def __init__(self, database_path: str | Path | None = None) -> None:
        self.database_path = Path(database_path or DATABASE_PATH)

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        return connection

    def get_counts(self, start_date: str, end_date: str) -> dict[str, int]:
        self._validate_dates(start_date, end_date)
        with self._connect() as con:
            tables = self._table_names(con)
            result = {
                "주문": self._count_period(con, "orders", "ordered_at", start_date, end_date),
                "입금": self._count_period(con, "payments", "paid_at", start_date, end_date),
                "배송": self._count_period(con, "shipments", "COALESCE(shipped_at, created_at)", start_date, end_date),
                "클레임": self._count_period(con, "claims", "created_at", start_date, end_date) if "claims" in tables else 0,
                "거래처": self._count_all(con, "customers"),
                "상품": self._count_all(con, "products"),
                "공급처": self._count_all(con, "suppliers"),
            }
        return result

    def create_backup(self, output_dir: str | Path) -> Path:
        destination = Path(output_dir)
        destination.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        target = destination / f"bisun_erp_backup_{stamp}.db"
        suffix = 1
        while target.exists():
            target = destination / f"bisun_erp_backup_{stamp}_{suffix}.db"
            suffix += 1
        shutil.copy2(self.database_path, target)
        return target

    def export_excel(
        self,
        output_dir: str | Path,
        start_date: str,
        end_date: str,
        selected: Iterable[str],
        *,
        create_backup: bool = True,
    ) -> tuple[Path, Path | None, dict[str, int]]:
        self._validate_dates(start_date, end_date)
        selected_set = set(selected)
        if not selected_set:
            raise ValueError("내보낼 데이터 항목을 하나 이상 선택하세요.")

        destination = Path(output_dir)
        destination.mkdir(parents=True, exist_ok=True)
        backup_path = self.create_backup(destination / "backup") if create_backup else None
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        excel_path = destination / f"비선상회_ERP_데이터_{start_date.replace('-', '')}_{end_date.replace('-', '')}_{stamp}.xlsx"

        workbook = Workbook()
        workbook.remove(workbook.active)
        counts: dict[str, int] = {}
        with self._connect() as con:
            tables = self._table_names(con)
            datasets = self._datasets(con, start_date, end_date, selected_set, tables)
            self._write_summary(workbook, start_date, end_date, datasets)
            for title, headers, rows in datasets:
                self._write_sheet(workbook, title, headers, rows)
                counts[title] = len(rows)
        workbook.save(excel_path)
        return excel_path, backup_path, counts

    def _datasets(self, con, start, end, selected, tables):
        result = []
        if "주문" in selected:
            result.append(("주문", ["플랫폼", "주문번호", "주문일", "수취인", "연락처", "주소", "주문상태", "결제상태", "매핑상태", "발주상태", "배송상태", "주문금액"], self._rows(con, """
                SELECT platform, order_number, ordered_at, receiver_name, receiver_phone,
                       TRIM(COALESCE(address,'') || ' ' || COALESCE(detail_address,'')),
                       order_status, payment_status, mapping_status, purchase_status,
                       shipment_status, total_amount
                FROM orders WHERE date(ordered_at) BETWEEN date(?) AND date(?)
                ORDER BY ordered_at, order_number
            """, start, end)))
        if "입금" in selected:
            result.append(("입금", ["주문번호", "입금자", "입금액", "입금방법", "입금상태", "입금일", "은행문자", "메모"], self._rows(con, """
                SELECT COALESCE(o.order_number,''), p.depositor_name, p.payment_amount,
                       p.payment_method, p.payment_status, p.paid_at, p.bank_message, p.memo
                FROM payments p LEFT JOIN orders o ON o.id=p.order_id
                WHERE date(COALESCE(p.paid_at,p.created_at)) BETWEEN date(?) AND date(?)
                ORDER BY COALESCE(p.paid_at,p.created_at)
            """, start, end)))
        if "배송" in selected:
            result.append(("배송", ["주문번호", "수취인", "택배사", "송장번호", "배송상태", "출고일", "완료일", "메모"], self._rows(con, """
                SELECT COALESCE(o.order_number,''), COALESCE(o.receiver_name,''), s.courier_name,
                       s.tracking_number, s.shipment_status, s.shipped_at, s.delivered_at, s.memo
                FROM shipments s LEFT JOIN orders o ON o.id=s.order_id
                WHERE date(COALESCE(s.shipped_at,s.created_at)) BETWEEN date(?) AND date(?)
                ORDER BY COALESCE(s.shipped_at,s.created_at)
            """, start, end)))
        if "클레임" in selected and "claims" in tables:
            columns = [row[1] for row in con.execute("PRAGMA table_info(claims)")]
            wanted = [c for c in ("id","order_id","claim_type","status","refund_status","refund_amount","reason","manager","memo","created_at","completed_at") if c in columns]
            rows = self._rows(con, f"SELECT {','.join(wanted)} FROM claims WHERE date(created_at) BETWEEN date(?) AND date(?) ORDER BY created_at", start, end)
            result.append(("클레임", wanted, rows))
        if "거래처" in selected:
            result.append(("거래처", ["유형", "상호명", "고객명", "연락처", "이메일", "사업자번호", "대표자", "주소", "메모"], self._rows(con, "SELECT customer_type,business_name,customer_name,phone,email,business_number,representative_name,TRIM(COALESCE(address,'')||' '||COALESCE(detail_address,'')),memo FROM customers ORDER BY business_name,customer_name")))
        if "상품" in selected:
            result.append(("상품", ["상품코드", "플랫폼", "플랫폼상품명", "상품명", "옵션", "공급처", "매입가", "판매가", "활성"], self._rows(con, "SELECT p.product_code,p.platform,p.platform_product_name,p.product_name,p.option_name,COALESCE(s.supplier_name,''),p.purchase_price,p.sale_price,CASE WHEN p.is_active=1 THEN '활성' ELSE '비활성' END FROM products p LEFT JOIN suppliers s ON s.id=p.supplier_id ORDER BY p.product_name")))
        if "공급처" in selected:
            result.append(("공급처", ["공급처코드", "공급처명", "담당자", "연락처", "이메일", "은행", "계좌번호", "예금주", "주소", "활성"], self._rows(con, "SELECT supplier_code,supplier_name,contact_name,phone,email,bank_name,bank_account,account_holder,address,CASE WHEN is_active=1 THEN '활성' ELSE '비활성' END FROM suppliers ORDER BY supplier_name")))
        return result

    @staticmethod
    def _rows(con, sql, *params):
        return [tuple(row) for row in con.execute(sql, params).fetchall()]

    @staticmethod
    def _write_summary(workbook, start, end, datasets):
        ws = workbook.create_sheet("요약")
        ws.append(["비선상회 ERP 데이터 내보내기"])
        ws.append(["조회기간", f"{start} ~ {end}"])
        ws.append(["생성일시", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        ws.append([])
        ws.append(["구분", "건수"])
        for title, _headers, rows in datasets:
            ws.append([title, len(rows)])
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 32
        ws["A1"].font = Font(size=16, bold=True)
        ws["A5"].font = ws["B5"].font = Font(bold=True)

    @staticmethod
    def _write_sheet(workbook, title, headers, rows):
        ws = workbook.create_sheet(title[:31])
        ws.append(headers)
        for row in rows:
            ws.append(list(row))
        fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for index, header in enumerate(headers, start=1):
            max_length = len(str(header))
            for cell in ws[get_column_letter(index)][1:201]:
                max_length = max(max_length, len(str(cell.value or "")))
            ws.column_dimensions[get_column_letter(index)].width = min(max_length + 3, 45)

    @staticmethod
    def _validate_dates(start, end):
        try:
            start_dt = datetime.strptime(start, "%Y-%m-%d")
            end_dt = datetime.strptime(end, "%Y-%m-%d")
        except ValueError as error:
            raise ValueError("날짜는 YYYY-MM-DD 형식으로 입력하세요.") from error
        if start_dt > end_dt:
            raise ValueError("시작일은 종료일보다 늦을 수 없습니다.")

    @staticmethod
    def _table_names(con):
        return {row[0] for row in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}

    @staticmethod
    def _count_all(con, table):
        return int(con.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0])

    @staticmethod
    def _count_period(con, table, field, start, end):
        return int(con.execute(f"SELECT COUNT(*) FROM {table} WHERE date({field}) BETWEEN date(?) AND date(?)", (start, end)).fetchone()[0])
