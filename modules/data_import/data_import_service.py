from __future__ import annotations

import re
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
DATABASE_PATH = PROJECT_ROOT / "data" / "bisun_erp.db"


class DataImportService:
    SUPPLIER_HEADERS = ["공급처명", "담당자", "연락처", "이메일", "발주방법", "기본배송비", "메모", "사용여부"]
    PRODUCT_HEADERS = ["상품코드", "판매처", "판매처상품명", "내부상품명", "옵션명", "공급처명", "공급처상품명", "발주단위", "매입단가", "판매가", "상품별마감시간", "사용여부"]

    def __init__(self, database_path: str | Path | None = None) -> None:
        self.database_path = Path(database_path or DATABASE_PATH)
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        self.ensure_schema()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.database_path)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys = ON")
        return conn

    def ensure_schema(self) -> None:
        with self._connect() as conn:
            supplier_columns = {r[1] for r in conn.execute("PRAGMA table_info(suppliers)")}
            product_columns = {r[1] for r in conn.execute("PRAGMA table_info(products)")}
            for name, sql_type in (("order_method", "TEXT"), ("default_shipping_fee", "INTEGER NOT NULL DEFAULT 0")):
                if name not in supplier_columns:
                    conn.execute(f"ALTER TABLE suppliers ADD COLUMN {name} {sql_type}")
            for name, sql_type in (("purchase_unit", "TEXT"), ("purchase_deadline", "TEXT")):
                if name not in product_columns:
                    conn.execute(f"ALTER TABLE products ADD COLUMN {name} {sql_type}")
            conn.commit()

    @staticmethod
    def _clean(value: Any) -> str:
        return "" if value is None else str(value).strip()

    @staticmethod
    def _bool(value: Any) -> int:
        text = str(value or "사용").strip().lower()
        return 0 if text in {"미사용", "비활성", "n", "no", "0", "false"} else 1

    @staticmethod
    def _number(value: Any, field: str) -> int:
        if value in (None, ""):
            return 0
        try:
            number = int(float(str(value).replace(",", "").strip()))
        except ValueError as exc:
            raise ValueError(f"{field}은 숫자로 입력해야 합니다.") from exc
        if number < 0:
            raise ValueError(f"{field}은 0 이상이어야 합니다.")
        return number

    @staticmethod
    def _deadline(value: Any) -> str:
        text = DataImportService._clean(value)
        if not text:
            return ""
        if re.fullmatch(r"\d{1,2}:\d{2}", text):
            hour, minute = map(int, text.split(":"))
            if 0 <= hour <= 23 and 0 <= minute <= 59:
                return f"{hour:02d}:{minute:02d}"
        raise ValueError("상품별마감시간은 14:30 형식으로 입력해야 합니다.")

    def create_template(self, path: str | Path) -> Path:
        path = Path(path)
        wb = Workbook()
        ws_s = wb.active
        ws_s.title = "공급처"
        ws_p = wb.create_sheet("상품")
        self._format_sheet(ws_s, self.SUPPLIER_HEADERS)
        self._format_sheet(ws_p, self.PRODUCT_HEADERS)
        ws_s.append(["예시수산", "홍길동", "010-0000-0000", "", "문자", 5000, "", "사용"])
        ws_p.append(["", "쿠팡", "자연산 전복 1kg", "자연산 전복", "1kg", "예시수산", "전복 1kg", "팩", 12000, 19000, "14:30", "사용"])
        ws_s.freeze_panes = "A2"; ws_p.freeze_panes = "A2"
        wb.save(path)
        return path

    @staticmethod
    def _format_sheet(ws, headers: list[str]) -> None:
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center")
        for idx, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(idx)].width = max(13, len(header) * 2 + 3)

    def validate_file(self, path: str | Path) -> dict[str, Any]:
        wb = load_workbook(path, data_only=True)
        errors: list[dict[str, Any]] = []
        suppliers: list[dict[str, Any]] = []
        products: list[dict[str, Any]] = []
        if "공급처" not in wb.sheetnames or "상품" not in wb.sheetnames:
            raise ValueError("엑셀 파일에 '공급처'와 '상품' 시트가 모두 필요합니다.")
        suppliers = self._read_suppliers(wb["공급처"], errors)
        products = self._read_products(wb["상품"], errors)
        supplier_names = {s["supplier_name"] for s in suppliers}
        with self._connect() as conn:
            supplier_names |= {r[0] for r in conn.execute("SELECT supplier_name FROM suppliers")}
            existing_product_keys = {(self._clean(r[0]), self._clean(r[1]), self._clean(r[2])) for r in conn.execute("SELECT product_name, option_name, supplier_id FROM products")}
        for item in products:
            if item["supplier_name"] and item["supplier_name"] not in supplier_names:
                errors.append({"sheet": "상품", "row": item["row"], "message": f"공급처 '{item['supplier_name']}'가 없습니다."})
        return {"suppliers": suppliers, "products": products, "errors": errors, "valid_count": len(suppliers)+len(products)-len(errors)}

    def _read_suppliers(self, ws, errors):
        headers = [self._clean(c.value) for c in ws[1]]
        missing = [h for h in self.SUPPLIER_HEADERS if h not in headers]
        if missing: raise ValueError(f"공급처 시트 필수 열 누락: {', '.join(missing)}")
        result=[]; names=set()
        for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True),2):
            if not any(v not in (None,"") for v in values): continue
            data=dict(zip(headers,values)); name=self._clean(data.get("공급처명"))
            try:
                if not name: raise ValueError("공급처명은 필수입니다.")
                if name in names: raise ValueError("파일 안에 같은 공급처명이 중복되었습니다.")
                names.add(name)
                result.append({"row":row_no,"supplier_name":name,"contact_name":self._clean(data.get("담당자")),"phone":self._clean(data.get("연락처")),"email":self._clean(data.get("이메일")),"order_method":self._clean(data.get("발주방법")),"shipping_fee":self._number(data.get("기본배송비"),"기본배송비"),"memo":self._clean(data.get("메모")),"is_active":self._bool(data.get("사용여부"))})
            except ValueError as e: errors.append({"sheet":"공급처","row":row_no,"message":str(e)})
        return result

    def _read_products(self, ws, errors):
        headers=[self._clean(c.value) for c in ws[1]]
        missing=[h for h in self.PRODUCT_HEADERS if h not in headers]
        if missing: raise ValueError(f"상품 시트 필수 열 누락: {', '.join(missing)}")
        result=[]; keys=set()
        for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True),2):
            if not any(v not in (None,"") for v in values): continue
            data=dict(zip(headers,values)); name=self._clean(data.get("내부상품명")); option=self._clean(data.get("옵션명")); supplier=self._clean(data.get("공급처명"))
            try:
                if not name: raise ValueError("내부상품명은 필수입니다.")
                if not supplier: raise ValueError("공급처명은 필수입니다.")
                if not self._clean(data.get("발주단위")): raise ValueError("발주단위는 필수입니다.")
                key=(name,option,supplier)
                if key in keys: raise ValueError("같은 상품·옵션·공급처 조합이 중복되었습니다.")
                keys.add(key)
                result.append({"row":row_no,"product_code":self._clean(data.get("상품코드")),"platform":self._clean(data.get("판매처")),"platform_product_name":self._clean(data.get("판매처상품명")),"product_name":name,"option_name":option,"supplier_name":supplier,"supplier_product_name":self._clean(data.get("공급처상품명")),"purchase_unit":self._clean(data.get("발주단위")),"purchase_price":self._number(data.get("매입단가"),"매입단가"),"sale_price":self._number(data.get("판매가"),"판매가"),"purchase_deadline":self._deadline(data.get("상품별마감시간")),"is_active":self._bool(data.get("사용여부"))})
            except ValueError as e: errors.append({"sheet":"상품","row":row_no,"message":str(e)})
        return result

    def import_file(self, path: str | Path) -> dict[str, Any]:
        result=self.validate_file(path)
        invalid_rows={(e["sheet"],e["row"]) for e in result["errors"]}
        created_s=updated_s=created_p=updated_p=0
        with self._connect() as conn:
            for s in result["suppliers"]:
                if ("공급처",s["row"]) in invalid_rows: continue
                row=conn.execute("SELECT id FROM suppliers WHERE supplier_name=?",(s["supplier_name"],)).fetchone()
                if row:
                    conn.execute("UPDATE suppliers SET contact_name=?,phone=?,email=?,order_method=?,default_shipping_fee=?,memo=?,is_active=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(s["contact_name"],s["phone"],s["email"],s["order_method"],s["shipping_fee"],s["memo"],s["is_active"],row[0])); updated_s+=1
                else:
                    code=f"SUP-{int(conn.execute('SELECT COALESCE(MAX(id),0)+1 FROM suppliers').fetchone()[0]):06d}"
                    conn.execute("INSERT INTO suppliers(supplier_code,supplier_name,contact_name,phone,email,memo,is_active,order_method,default_shipping_fee,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)",(code,s["supplier_name"],s["contact_name"],s["phone"],s["email"],s["memo"],s["is_active"],s["order_method"],s["shipping_fee"])); created_s+=1
            supplier_map={r["supplier_name"]:r["id"] for r in conn.execute("SELECT id,supplier_name FROM suppliers")}
            for p in result["products"]:
                if ("상품",p["row"]) in invalid_rows: continue
                sid=supplier_map[p["supplier_name"]]
                row=conn.execute("SELECT id FROM products WHERE product_name=? AND COALESCE(option_name,'')=? AND supplier_id=?",(p["product_name"],p["option_name"],sid)).fetchone()
                if row:
                    conn.execute("UPDATE products SET platform=?,platform_product_name=?,supplier_product_name=?,purchase_price=?,sale_price=?,purchase_unit=?,purchase_deadline=?,is_active=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(p["platform"],p["platform_product_name"],p["supplier_product_name"],p["purchase_price"],p["sale_price"],p["purchase_unit"],p["purchase_deadline"],p["is_active"],row[0])); updated_p+=1
                else:
                    code=p["product_code"] or f"PRD-{int(conn.execute('SELECT COALESCE(MAX(id),0)+1 FROM products').fetchone()[0]):06d}"
                    conn.execute("INSERT INTO products(product_code,platform,platform_product_name,product_name,option_name,supplier_id,supplier_product_name,purchase_price,sale_price,purchase_unit,purchase_deadline,is_active,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)",(code,p["platform"],p["platform_product_name"],p["product_name"],p["option_name"],sid,p["supplier_product_name"],p["purchase_price"],p["sale_price"],p["purchase_unit"],p["purchase_deadline"],p["is_active"])); created_p+=1
            conn.commit()
        return {"created_suppliers":created_s,"updated_suppliers":updated_s,"created_products":created_p,"updated_products":updated_p,"errors":result["errors"]}
