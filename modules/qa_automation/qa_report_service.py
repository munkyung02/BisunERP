from __future__ import annotations

import json
from datetime import datetime, timedelta
from pathlib import Path
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.qa.models import QAReport


PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
OUTPUT_ROOT = PROJECT_ROOT / "output" / "qa_reports"


class QAReportService:
    def __init__(
        self,
        output_root: str | Path = OUTPUT_ROOT,
        now_provider: Callable[[], datetime] | None = None,
    ) -> None:
        self.output_root = Path(output_root)
        self.now_provider = now_provider or datetime.now

    def export_excel(self, report: QAReport) -> str:
        path = self._output_path("xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        summary = workbook.active
        summary.title = "QA Summary"
        summary.append(["Bisun ERP QA Automation Report"])
        summary["A1"].font = Font(size=16, bold=True)
        summary.merge_cells("A1:E1")
        row_number = 3
        for key, value in report.metadata.items():
            summary.cell(row=row_number, column=1, value=key).font = Font(bold=True)
            summary.cell(row=row_number, column=2, value=str(value)).number_format = "@"
            row_number += 1
        row_number += 1
        summary.cell(row=row_number, column=1, value="PASS")
        summary.cell(row=row_number, column=2, value=report.pass_count)
        summary.cell(row=row_number, column=3, value="WARNING")
        summary.cell(row=row_number, column=4, value=report.warning_count)
        summary.cell(row=row_number, column=5, value="FAILED")
        summary.cell(row=row_number, column=6, value=report.failed_count)
        row_number += 2
        headers = ["No", "Check", "Status", "Elapsed", "Summary"]
        summary.append(headers)
        header_row = summary.max_row
        for cell in summary[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="355E4B")
        for index, result in enumerate(report.results, start=1):
            summary.append([
                index,
                result.name,
                result.status,
                result.elapsed_seconds,
                result.summary,
            ])
        summary.append([])
        summary.append(["Total Elapsed", report.elapsed_seconds])
        for column, width in enumerate((8, 32, 13, 13, 70), start=1):
            summary.column_dimensions[get_column_letter(column)].width = width
        for row in summary.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        details = workbook.create_sheet("QA Details")
        details.append(["Check ID", "Check", "Status", "Elapsed", "Detail", "Data"])
        for cell in details[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="355E4B")
        for result in report.results:
            details.append([
                result.check_id,
                result.name,
                result.status,
                result.elapsed_seconds,
                result.detail,
                json.dumps(result.data, ensure_ascii=False, default=str),
            ])
        for column, width in enumerate((24, 32, 13, 13, 80, 80), start=1):
            details.column_dimensions[get_column_letter(column)].width = width
        for row in details.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        workbook.save(path)
        workbook.close()
        return str(path)

    def export_pdf(self, report: QAReport) -> str:
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_LEFT
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        except ImportError as error:
            raise RuntimeError("PDF 생성을 위해 reportlab 패키지가 필요합니다.") from error

        path = self._output_path("pdf")
        path.parent.mkdir(parents=True, exist_ok=True)
        pdfmetrics.registerFont(UnicodeCIDFont("HYSMyeongJo-Medium"))
        styles = getSampleStyleSheet()
        normal = ParagraphStyle(
            "KoreanNormal",
            parent=styles["Normal"],
            fontName="HYSMyeongJo-Medium",
            fontSize=8,
            leading=11,
            alignment=TA_LEFT,
        )
        title = ParagraphStyle(
            "KoreanTitle",
            parent=normal,
            fontSize=16,
            leading=20,
            spaceAfter=8,
        )
        document = SimpleDocTemplate(
            str(path),
            pagesize=landscape(A4),
            leftMargin=12 * mm,
            rightMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
            title="Bisun ERP QA Automation Report",
        )
        story = [Paragraph("Bisun ERP QA Automation Report", title)]
        metadata_rows = [
            [Paragraph(str(key), normal), Paragraph(str(value), normal)]
            for key, value in report.metadata.items()
        ]
        metadata_table = Table(metadata_rows, colWidths=[48 * mm, 205 * mm])
        metadata_table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E6EFE9")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.extend([
            metadata_table,
            Spacer(1, 7 * mm),
            Paragraph(
                f"PASS {report.pass_count} · WARNING {report.warning_count} · "
                f"FAILED {report.failed_count} · Total {report.elapsed_seconds:.2f}s",
                normal,
            ),
            Spacer(1, 3 * mm),
        ])
        result_rows = [[
            Paragraph("No", normal),
            Paragraph("Check", normal),
            Paragraph("Status", normal),
            Paragraph("Elapsed", normal),
            Paragraph("Summary", normal),
        ]]
        for index, result in enumerate(report.results, start=1):
            result_rows.append([
                index,
                Paragraph(result.name, normal),
                result.status,
                f"{result.elapsed_seconds:.2f}s",
                Paragraph(result.summary, normal),
            ])
        result_table = Table(result_rows, colWidths=[10 * mm, 55 * mm, 24 * mm, 22 * mm, 142 * mm], repeatRows=1)
        result_table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#355E4B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.extend([result_table, PageBreak(), Paragraph("Detailed Results", title)])
        for result in report.results:
            story.extend([
                Paragraph(f"{result.name} · {result.status}", normal),
                Paragraph(result.detail or result.summary, normal),
                Spacer(1, 3 * mm),
            ])
        document.build(story)
        return str(path)

    def _output_path(self, extension: str) -> Path:
        current = self.now_provider()
        while True:
            directory = self.output_root / current.strftime("%Y%m%d")
            path = directory / f"QA_Report_{current.strftime('%Y%m%d_%H%M%S')}.{extension}"
            if not path.exists():
                return path
            current += timedelta(seconds=1)
