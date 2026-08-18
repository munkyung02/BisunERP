from __future__ import annotations

import re
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .base import PurchaseTemplate
from .helpers import full_address, get_unique_output_path, safe_filename
from modules.products.product_shipping_policy_repository import (
    ProductShippingPolicyRepository,
)


PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
DATABASE_PATH = PROJECT_ROOT / "data" / "bisun_erp.db"


class FoodPresidentPurchaseTemplate(PurchaseTemplate):
    """푸드대통령 전용 발주서 템플릿

    - 컬럼 A-P(16열)를 사용합니다.
    - 주문번호 형식: YYMMDDNN (일별 연속 증가)
    - 공급단가1(열 O): unit_price * purchase_quantity + shipping_fee
      shipping_fee는 `ProductShippingPolicyRepository.get_shipping_fee()`에서 조회합니다.
      조회 결과가 `None`이거나 복수 정책(ValueError)이면 즉시 예외를 발생시켜 발주 중단합니다.
    """

    template_key = "foodpresident"
    display_name = "푸드대통령 전용 발주서"

    def create_excel(
        self,
        *,
        output_directory: Path,
        supplier_name: str,
        purchase_round: str,
        items: list[dict[str, Any]],
    ) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"

        headers = [
            "모델명",  # A
            "상품주문번호",  # B
            "발주일자",  # C
            "상품명",  # D
            "수량",  # E
            "주문자",  # F
            "주문자연락처1",  # G
            "주문자연락처2",  # H
            "수령인",  # I
            "수령인연락처1",  # J
            "수령인연락처2",  # K
            "우편번호",  # L
            "주소",  # M
            "배송메모",  # N
            "공급단가1",  # O
            "공급단가2(O열복사)",  # P (exact required header)
        ]

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        header_fill = PatternFill(fill_type="solid", fgColor="DDDDDD")

        # place header at row 1 per original layout
        header_row = 1
        for col_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=header_row, column=col_index, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            cell.fill = header_fill
        # determine next sequence for today's order numbers using existing purchase_orders
        today_prefix = datetime.now().strftime("%y%m%d")
        conn = sqlite3.connect(DATABASE_PATH)
        try:
            conn.row_factory = sqlite3.Row
            rows = conn.execute(
                "SELECT order_number FROM purchase_orders WHERE order_number LIKE ?",
                (f"{today_prefix}%",),
            ).fetchall()
        finally:
            conn.close()

        existing_seqs: list[int] = []
        for r in rows:
            on = str(r["order_number"] or "").strip()
            m = re.fullmatch(rf"{today_prefix}(\d{{2}})$", on)
            if m:
                try:
                    existing_seqs.append(int(m.group(1)))
                except Exception:
                    pass

        next_seq = max(existing_seqs) + 1 if existing_seqs else 1

        shipping_repo = ProductShippingPolicyRepository()

        # data rows start at row 2
        row_index = 2
        for item in items:
            product_id = item.get("product_id")
            if not product_id:
                raise ValueError("템플릿 항목에 product_id가 없습니다.")

            qty = int(item.get("purchase_quantity", item.get("quantity", 0)) or 0)
            if qty <= 0:
                raise ValueError(f"product_id={product_id}: 발주수량이 0이거나 올바르지 않습니다.")

            unit_price = int(item.get("unit_price") or 0)

            # get shipping fee; may raise ValueError for overlapping policies
            shipping_fee = shipping_repo.get_shipping_fee(int(product_id), int(qty))
            if shipping_fee is None:
                raise ValueError(
                    f"product_id={product_id} 수량={qty}에 대한 배송비 정책이 없습니다. Notion에서 배송정책을 확인하세요."
                )

            total_product_amount = unit_price * qty
            supply_price1 = total_product_amount + int(shipping_fee)

            order_number = f"{today_prefix}{next_seq:02d}"
            next_seq += 1

            supplier_product_name = (
                item.get("supplier_product_name")
                or item.get("product_name")
                or item.get("platform_product_name")
                or ""
            )

            # Orderer fields: check common keys; do NOT substitute with receiver_name
            orderer_name = None
            orderer_phone1 = None
            for key in ("orderer_name", "buyer_name", "주문자명", "orderer", "buyer"):
                if key in item and item.get(key) not in (None, ""):
                    orderer_name = item.get(key)
                    break
            for key in ("orderer_phone", "buyer_phone", "주문자_휴대폰", "orderer_phone1", "buyer_phone1"):
                if key in item and item.get(key) not in (None, ""):
                    orderer_phone1 = item.get(key)
                    break

            receiver_name = item.get("receiver_name") or ""
            receiver_phone = item.get("receiver_phone") or ""
            postal_code = item.get("postal_code") or ""
            address = full_address(item.get("address"), item.get("detail_address")) or ""

            # 발주일자: use today as date object to preserve Excel date format
            order_date = datetime.now().date()

            # If orderer_name missing, leave F/G blank (do not substitute)
            f_value = orderer_name or ""
            g_value = orderer_phone1 or ""

            # Determine delivery_message: prefer item value, else try DB read-only lookup by order_item_id
            delivery_note = item.get("delivery_message")
            if not delivery_note:
                # attempt read-only lookup
                try:
                    db_uri = f"file:///{DATABASE_PATH.as_posix()}?mode=ro"
                    conn_lookup = sqlite3.connect(db_uri, uri=True)
                    conn_lookup.row_factory = sqlite3.Row
                    cur_lookup = conn_lookup.cursor()
                    oid = item.get("order_item_id")
                    if oid is not None:
                        row_dm = cur_lookup.execute(
                            "SELECT delivery_message FROM purchase_orders WHERE order_item_id = ?",
                            (int(oid),),
                        ).fetchone()
                        if row_dm is not None:
                            delivery_note = row_dm["delivery_message"]
                    conn_lookup.close()
                except Exception:
                    delivery_note = delivery_note or None

            # If delivery_note is missing in item and DB, leave N blank per spec
            n_value = delivery_note or ""

            values = [
                "",  # A 모델명 (blank per original)
                order_number,  # B 상품주문번호
                order_date,  # C 발주일자
                supplier_product_name,  # D 상품명
                int(qty),  # E 수량
                f_value,  # F 주문자
                g_value,  # G 주문자연락처1
                "",  # H 주문자연락처2
                receiver_name,  # I 수령인
                receiver_phone,  # J 수령인연락처1
                "",  # K 수령인연락처2
                postal_code,  # L 우편번호
                address,  # M 주소
                n_value,  # N 배송메모
                int(supply_price1),  # O 공급단가1
                int(supply_price1),  # P 공급단가2 (copy)
            ]

            for col_idx, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_index, column=col_idx, value=value)
                cell.border = thin_border
                if col_idx in {15, 16}:
                    cell.number_format = '#,##0"원"'
                if col_idx == 3 and isinstance(value, (str,)):
                    # leave as-is
                    pass
                if col_idx == 3 and hasattr(value, 'strftime'):
                    cell.number_format = 'yyyy-mm-dd'
                cell.alignment = Alignment(vertical="center")

            row_index += 1

        created_date = datetime.now().strftime("%Y%m%d")
        file_path = get_unique_output_path(
            output_directory / f"푸드대통령_{created_date} 더유_발주서.xlsx"
        )

        workbook.save(file_path)

        return file_path
