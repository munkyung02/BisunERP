from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .base import PurchaseTemplate
from .helpers import full_address, get_unique_output_path


class MagoPurchaseTemplate(PurchaseTemplate):
    """주식회사 마고 전용 택배 발주서 템플릿."""

    template_key = "mago"
    display_name = "주식회사 마고 발주서"

    TEMPLATE_PATH = (
        Path(__file__).resolve().parent
        / "assets"
        / "mago_delivery_template.xlsx"
    )

    SENDER_NAME = "비선상회"
    SENDER_PHONE = "010-2431-0204"
    SENDER_ADDRESS = ""
    COMPANY_NAME = "주식회사 마고"

    def create_excel(
        self,
        *,
        output_directory: Path,
        supplier_name: str,
        purchase_round: str,
        items: list[dict[str, Any]],
    ) -> Path:
        if not self.TEMPLATE_PATH.is_file():
            raise FileNotFoundError(
                f"마고 발주서 원본 템플릿이 없습니다: {self.TEMPLATE_PATH}"
            )

        workbook = load_workbook(self.TEMPLATE_PATH)
        if workbook.sheetnames != ["Sheet1"]:
            workbook.close()
            raise ValueError("마고 발주서 원본의 시트 구성이 변경되었습니다.")

        sheet = workbook["Sheet1"]
        expected_headers = [
            "보내시는분",
            "보내시는분 전화",
            "보내는분 총주소",
            "받으시는 분",
            "받으시는 분 전화",
            "받는분 총 주소",
            "물품명(무게)",
            "수량",
            "배송메모",
            "업체명",
            "    ",
        ]
        actual_headers = [sheet.cell(1, column).value for column in range(1, 12)]
        if actual_headers != expected_headers:
            workbook.close()
            raise ValueError("마고 발주서 원본의 A1:K1 헤더가 변경되었습니다.")

        # 원본에는 857행까지 빈 셀 스타일이 있습니다. 실제 발주 행만 남기되,
        # 데이터 행의 원본 스타일은 그대로 복사해 사용합니다.
        template_row = 2
        required_rows = max(len(items), 1)
        for row_index in range(2, 2 + required_rows):
            if row_index > sheet.max_row:
                sheet.insert_rows(row_index)
                for column in range(1, 12):
                    source = sheet.cell(template_row, column)
                    target = sheet.cell(row_index, column)
                    target._style = copy(source._style)
                    target.number_format = source.number_format

        for row_index, item in enumerate(items, start=2):
            product_name = (
                item.get("supplier_product_name")
                or item.get("product_name")
                or item.get("platform_product_name")
                or ""
            )
            quantity = int(item.get("purchase_quantity") or 0)
            values = [
                self.SENDER_NAME,
                self.SENDER_PHONE,
                self.SENDER_ADDRESS,
                item.get("receiver_name") or "",
                str(item.get("receiver_phone") or ""),
                full_address(item.get("address"), item.get("detail_address")),
                product_name,
                quantity,
                item.get("delivery_message") or "",
                self.COMPANY_NAME,
                None,
            ]
            for column, value in enumerate(values, start=1):
                sheet.cell(row_index, column, value=value)

            # 발신/수신 전화번호의 선행 0을 보존합니다.
            sheet.cell(row_index, 2).number_format = "@"
            sheet.cell(row_index, 5).number_format = "@"

        first_unused_row = 2 + len(items)
        if first_unused_row <= sheet.max_row:
            sheet.delete_rows(
                first_unused_row,
                sheet.max_row - first_unused_row + 1,
            )

        output_directory.mkdir(parents=True, exist_ok=True)
        created_date = datetime.now().strftime("%Y%m%d")
        file_path = get_unique_output_path(
            output_directory
            / f"{created_date}_상품출고요청서_더유_주식회사마고.xlsx"
        )
        workbook.save(file_path)
        workbook.close()
        return file_path
