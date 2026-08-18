from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .base import PurchaseTemplate
from .helpers import full_address, get_unique_output_path, safe_filename


class MasterYutongPurchaseTemplate(PurchaseTemplate):
    """마스터유통 전용 발주서 템플릿"""

    template_key = "masteryutong"
    display_name = "마스터유통 발주서"

    FIRST_ROW = (
        "상품명을 임의로 작성해 주실 경우 어떤 상품인지 몰라 상품이 잘못 출고 될 수 있습니다. 되도록 시트 내 상품명으로 발주 부탁드립니다."
    )
    SECOND_ROW = (
        "A = 업체명 기입해 주세요",
        "B = 주소가 필요 없으시면 안적어 주셔도 됩니다.",
        "C = 업체 연락처 기입해 주세요",
        "F = 상품명은 상품명 x 'n' 입니다.",
    )

    HEADERS = [
        "주문처",
        "주문처 주소",
        "주문처 연락처",
        "공급사",
        "창고명",
        "상품명",
        "성함",
        "주소",
        "연락처",
        "배송메세지/특이사항",
    ]

    def create_excel(
        self,
        *,
        output_directory: Path,
        supplier_name: str,
        purchase_round: str,
        items: list[dict[str, Any]],
    ) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"

        # Row1: full guidance merged across A:J
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
        cell = sheet.cell(row=1, column=1, value=self.FIRST_ROW)
        cell.alignment = Alignment(wrap_text=True)

        # Row2: multiple guidance pieces placed in some columns per spec
        sheet.cell(row=2, column=1, value=self.SECOND_ROW[0])
        sheet.cell(row=2, column=2, value=self.SECOND_ROW[1])
        sheet.cell(row=2, column=3, value=self.SECOND_ROW[2])
        sheet.cell(row=2, column=6, value=self.SECOND_ROW[3])

        # Row3: headers
        for col_index, header in enumerate(self.HEADERS, start=1):
            cell = sheet.cell(row=3, column=col_index, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data rows start at row 4
        for row_index, item in enumerate(items, start=4):
            # A = fixed orderer shop name
            sheet.cell(row=row_index, column=1, value="비선상회")
            # B = blank
            sheet.cell(row=row_index, column=2, value="")
            # C = fixed phone
            sheet.cell(row=row_index, column=3, value="010-2431-0204")
            # D = blank
            sheet.cell(row=row_index, column=4, value="")
            # E = blank
            sheet.cell(row=row_index, column=5, value="")

            # F = supplier_product_name x purchase_quantity
            product = item.get("supplier_product_name") or item.get("product_name") or ""
            pq = int(item.get("purchase_quantity") or item.get("quantity") or 0)
            f_value = f"{product} x {pq}"
            sheet.cell(row=row_index, column=6, value=f_value)

            # G = receiver_name
            sheet.cell(row=row_index, column=7, value=item.get("receiver_name") or "")
            # H = full address
            sheet.cell(row=row_index, column=8, value=full_address(item.get("address"), item.get("detail_address")) or "")
            # I = receiver_phone
            sheet.cell(row=row_index, column=9, value=item.get("receiver_phone") or "")
            # J = delivery_message
            sheet.cell(row=row_index, column=10, value=item.get("delivery_message") or "")

            sheet.row_dimensions[row_index].height = 20

        output_directory.mkdir(parents=True, exist_ok=True)
        created_date = datetime.now().strftime("%Y%m%d")
        file_name = f"{created_date}_더유_마스터.xlsx"
        file_path = get_unique_output_path(output_directory / file_name)
        workbook.save(file_path)
        return file_path
