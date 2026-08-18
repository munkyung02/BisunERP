from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .base import PurchaseTemplate
from .helpers import full_address, get_unique_output_path, safe_filename


class MangoB2BPuchaseTemplate(PurchaseTemplate):
    """망고B2B 전용 발주서 템플릿"""

    template_key = "mangob2b"
    display_name = "망고B2B 발주서"

    HEADERS = [
        "상호명",
        "상호 연락처",
        "주문상품명(옵션포함)",
        "수량",
        "수령인",
        "수령인 휴대전화",
        "수령인 주소",
        "수령인 우편번호",
        "배송메세지",
        "발주일",
        "택배사",
        "운송장번호",
    ]

    FIXED_SHOP = "비선상회"
    FIXED_SHOP_PHONE = "010-2431-0204"

    def create_excel(
        self,
        *,
        output_directory: Path,
        supplier_name: str,
        purchase_round: str,
        items: list[dict[str, Any]],
    ) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"

        # Write headers A1:L1
        for col_index, header in enumerate(self.HEADERS, start=1):
            cell = sheet.cell(row=1, column=col_index, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Populate rows from row 2
        for row_index, item in enumerate(items, start=2):
            # A: fixed shop name
            sheet.cell(row=row_index, column=1, value=self.FIXED_SHOP)
            # B: fixed shop contact
            sheet.cell(row=row_index, column=2, value=self.FIXED_SHOP_PHONE)

            # C: supplier_product_name or fallback to product_name; include option_name if present by concatenation
            product = item.get("supplier_product_name") or item.get("product_name") or ""
            option = item.get("option_name")
            if option and str(option).strip():
                product_cell = f"{product} {option}"
            else:
                product_cell = product
            sheet.cell(row=row_index, column=3, value=product_cell)

            # D: purchase_quantity (final)
            sheet.cell(row=row_index, column=4, value=int(item.get("purchase_quantity") or item.get("quantity") or 0))

            # E: receiver_name
            sheet.cell(row=row_index, column=5, value=item.get("receiver_name") or "")
            # F: receiver_phone
            sheet.cell(row=row_index, column=6, value=item.get("receiver_phone") or "")
            # G: full address
            sheet.cell(row=row_index, column=7, value=full_address(item.get("address"), item.get("detail_address")) or "")
            # H: postal_code
            sheet.cell(row=row_index, column=8, value=item.get("postal_code") or "")
            # I: delivery_message
            sheet.cell(row=row_index, column=9, value=item.get("delivery_message") or "")
            # J: 발주일 -> use purchased_at if present else created_at else empty
            sheet.cell(row=row_index, column=10, value=item.get("purchased_at") or item.get("created_at") or "")

            # K, L blank per spec
            sheet.cell(row=row_index, column=11, value="")
            sheet.cell(row=row_index, column=12, value="")

            sheet.row_dimensions[row_index].height = 20

        output_directory.mkdir(parents=True, exist_ok=True)
        created_date = datetime.now().strftime("%Y%m%d")
        file_name = f"{created_date}.mgb2bmall_더유.xlsx"
        file_path = get_unique_output_path(output_directory / file_name)
        workbook.save(file_path)
        return file_path
