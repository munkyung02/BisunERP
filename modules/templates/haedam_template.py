from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .base import PurchaseTemplate
from .helpers import full_address, safe_filename


class HaedamPurchaseTemplate(PurchaseTemplate):
    """해담 전용 발주서(출고요청서) 템플릿

    단일 시트(Sheet1), 헤더 row=1, 데이터 row=2부터
    컬럼 순서와 서식은 제공된 기준 파일을 최대한 일치시킵니다.
    """

    template_key = "haedam"
    display_name = "해담 전용 발주서"

    def create_excel(
        self,
        *,
        output_directory: Path,
        supplier_name: str,
        purchase_round: str,
        items: list[dict[str, Any]],
    ) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"

        # Header definitions (exact order)
        headers = [
            "보내는분성명",
            "보내는분전화번호",
            "보내는분주소(전체, 분할)",
            "받는분성명",
            "받는분전화번호",
            "받는분주소(전체, 분할)",
            "품목",
            "수량",
            "배송메세지",
            "송장번호",
        ]

        # Column widths (A-J)
        widths = {
            1: 14.19921875,
            2: 17.19921875,
            3: 48.3984375,
            4: 11.0,
            5: 16.5,
            6: 80.59765625,
            7: 47.09765625,
            8: 34.5,
            9: 9.0,
            10: 13.0,
        }

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        header_fill = PatternFill(fill_type="solid", fgColor="DDDDDD")
        header_font = Font(name="맑은 고딕", size=11, bold=False)
        header_alignment = Alignment(horizontal="left", vertical="center")

        # write headers on row 1
        for col_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=col_index, value=header)
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
            cell.fill = PatternFill(fill_type="solid", fgColor="FFDDDDDD")

        # column widths
        from openpyxl.utils import get_column_letter

        for col_index, width in widths.items():
            sheet.column_dimensions[get_column_letter(col_index)].width = width

        # fixed sender values
        sender_name = "더유"
        sender_phone = "010-2431-0204"

        # write data rows starting at row 2
        row = 2
        for item in items:
            receiver_name = item.get("receiver_name") or ""
            receiver_phone = item.get("receiver_phone") or ""
            address = full_address(item.get("address"), item.get("detail_address")) or ""

            supplier_product_name = (
                item.get("supplier_product_name")
                or item.get("product_name")
                or item.get("platform_product_name")
                or ""
            )

            # final quantity must use purchase_quantity when available
            # Use only the final calculated purchase_quantity.
            # Do not fallback to raw 'quantity' or re-calculate from options.
            quantity = int(item.get("purchase_quantity") or 0)

            delivery_message = item.get("delivery_message") or ""

            values = [
                sender_name,
                sender_phone,
                "",  # 보내는분주소(전체, 분할) - intentionally blank
                receiver_name,
                receiver_phone,
                address,
                supplier_product_name,
                int(quantity or 0),
                delivery_message,
                "",  # 송장번호 empty on generation
            ]

            for col_index, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=col_index, value=value)
                # basic alignment similar to sample
                if col_index in {1, 2, 4, 5, 8}:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(vertical="center")
            row += 1

        timestamp = datetime.now().strftime("%H%M%S")
        file_path = output_directory / (
            f"{safe_filename(supplier_name)}_{safe_filename(purchase_round)}_해담_출고요청서_{timestamp}.xlsx"
        )

        workbook.save(file_path)

        return file_path
