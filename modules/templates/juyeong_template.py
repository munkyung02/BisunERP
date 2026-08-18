from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .base import PurchaseTemplate
from .helpers import full_address, get_unique_output_path, safe_filename


class JuyeongSeafoodPurchaseTemplate(PurchaseTemplate):
    """주영씨푸드 전용 발주서 템플릿 (A:N 고정 레이아웃)"""

    template_key = "juyeong"
    display_name = "주영씨푸드 발주서"

    HEADERS = [
        "받는분성명",
        "전화번호",
        "받는분주소",
        "품목명",
        "박스수량",
        "박스타입",
        "배송메세지1",
        "운임구분",
        "보내는분성명",
        "보내는분전화번호",
        "보내는분 주소",
        "",
        "택배사",
        "운송장번호",
    ]

    def create_excel(
        self,
        *,
        output_directory: Path,
        supplier_name: str,
        purchase_round: str,
        items: list[dict[str, Any]],
    ) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"

        # Header A1:N1
        for col_index, header in enumerate(self.HEADERS, start=1):
            cell = sheet.cell(row=1, column=col_index, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data rows start at row 2
        for row_index, item in enumerate(items, start=2):
            # A: receiver_name
            sheet.cell(row=row_index, column=1, value=item.get("receiver_name") or "")
            # B: receiver_phone
            sheet.cell(row=row_index, column=2, value=item.get("receiver_phone") or "")
            # C: full address (address + detail_address)
            sheet.cell(
                row=row_index,
                column=3,
                value=full_address(item.get("address"), item.get("detail_address")) or "",
            )

            # D: supplier_product_name x purchase_quantity (fallback to product_name only if supplier_product_name empty)
            supplier_name_field = (
                item.get("supplier_product_name")
                or item.get("product_name")
                or ""
            )
            purchase_qty = item.get("purchase_quantity")
            if purchase_qty is None:
                purchase_qty = item.get("quantity")

            d_value = f"{supplier_name_field} x {int(purchase_qty or 0)}"
            sheet.cell(row=row_index, column=4, value=d_value)

            # E: 박스수량 always 1
            sheet.cell(row=row_index, column=5, value=1)

            # F: 박스타입 (blank per spec)
            sheet.cell(row=row_index, column=6, value="")

            # G: 배송메세지1
            sheet.cell(row=row_index, column=7, value=item.get("delivery_message") or "")

            # H: 운임구분 blank
            sheet.cell(row=row_index, column=8, value="")

            # I: fixed sender name
            sheet.cell(row=row_index, column=9, value="비선상회")

            # J: fixed sender phone
            sheet.cell(row=row_index, column=10, value="010-2431-0204")

            # K: 보내는분 주소 blank
            sheet.cell(row=row_index, column=11, value="")

            # L: header blank (keep empty)
            sheet.cell(row=row_index, column=12, value="")

            # M: 택배사 blank
            sheet.cell(row=row_index, column=13, value="")

            # N: 운송장번호 blank
            sheet.cell(row=row_index, column=14, value="")

            sheet.row_dimensions[row_index].height = 22

        # Save file
        output_directory.mkdir(parents=True, exist_ok=True)
        created_date = datetime.now().strftime("%y%m%d")
        file_name = f"{created_date}_상품출고요청서_더유_주영씨푸드.xlsx"
        file_path = get_unique_output_path(output_directory / file_name)

        workbook.save(file_path)
        return file_path
