from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .base import PurchaseTemplate
from .helpers import aggregate_items, full_address, safe_filename


class DefaultPurchaseTemplate(PurchaseTemplate):
    """
    비선ERP 기본 발주서입니다.

    시트:
    - 발주요약: 동일 상품 합산
    - 배송목록: 주문별 배송정보
    """

    template_key = "default"
    display_name = "기본 발주서"

    def create_excel(
        self,
        *,
        output_directory: Path,
        supplier_name: str,
        purchase_round: str,
        items: list[dict[str, Any]],
    ) -> Path:
        workbook = Workbook()

        summary_sheet = workbook.active
        summary_sheet.title = "발주요약"
        summary_sheet.sheet_view.showGridLines = False
        summary_sheet.freeze_panes = "A5"

        detail_sheet = workbook.create_sheet("배송목록")
        detail_sheet.sheet_view.showGridLines = False
        detail_sheet.freeze_panes = "A5"

        thin_border = Border(
            left=Side(style="thin", color="B7B7B7"),
            right=Side(style="thin", color="B7B7B7"),
            top=Side(style="thin", color="B7B7B7"),
            bottom=Side(style="thin", color="B7B7B7"),
        )
        header_fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAD3",
        )

        aggregated_items = aggregate_items(items)

        self._write_summary_sheet(
            sheet=summary_sheet,
            supplier_name=supplier_name,
            purchase_round=purchase_round,
            items=items,
            aggregated_items=aggregated_items,
            thin_border=thin_border,
            header_fill=header_fill,
        )

        self._write_detail_sheet(
            sheet=detail_sheet,
            supplier_name=supplier_name,
            purchase_round=purchase_round,
            items=items,
            thin_border=thin_border,
            header_fill=header_fill,
        )

        timestamp = datetime.now().strftime("%H%M%S")
        file_name = (
            f"{safe_filename(supplier_name)}_"
            f"{safe_filename(purchase_round)}_"
            f"발주서_{timestamp}.xlsx"
        )
        file_path = output_directory / file_name

        workbook.active = 0
        workbook.save(file_path)

        return file_path

    @staticmethod
    def _write_summary_sheet(
        *,
        sheet,
        supplier_name: str,
        purchase_round: str,
        items: list[dict[str, Any]],
        aggregated_items: list[dict[str, Any]],
        thin_border: Border,
        header_fill: PatternFill,
    ) -> None:
        sheet.merge_cells("A1:H1")
        sheet["A1"] = "비선상회 공급처 발주요약"
        sheet["A1"].font = Font(size=18, bold=True)
        sheet["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        sheet.row_dimensions[1].height = 32

        sheet["A2"] = "공급처"
        sheet["B2"] = supplier_name
        sheet["D2"] = "발주차수"
        sheet["E2"] = purchase_round
        sheet["G2"] = "생성일시"
        sheet["H2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        total_quantity = sum(
            int(item.get("total_quantity") or 0)
            for item in aggregated_items
        )
        total_amount = sum(
            int(item.get("total_amount") or 0)
            for item in aggregated_items
        )

        sheet["A3"] = "상품 종류"
        sheet["B3"] = len(aggregated_items)
        sheet["C3"] = "주문상품 건수"
        sheet["D3"] = len(items)
        sheet["E3"] = "총 수량"
        sheet["F3"] = total_quantity
        sheet["G3"] = "총 상품금액"
        sheet["H3"] = total_amount
        sheet["H3"].number_format = '#,##0"원"'

        headers = [
            "번호",
            "공급처 상품코드",
            "공급처 상품명",
            "옵션",
            "주문건수",
            "총 수량",
            "매입단가",
            "총 상품금액",
        ]

        for column_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=4, column=column_index, value=header)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row_index, item in enumerate(aggregated_items, start=1):
            excel_row = 4 + row_index
            values = [
                row_index,
                item.get("supplier_product_code") or "",
                item.get("product_name") or "",
                item.get("option_name") or "",
                int(item.get("order_count") or 0),
                int(item.get("total_quantity") or 0),
                int(item.get("unit_price") or 0),
                int(item.get("total_amount") or 0),
            ]

            for column_index, value in enumerate(values, start=1):
                cell = sheet.cell(
                    row=excel_row,
                    column=column_index,
                    value=value,
                )
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal=(
                        "left"
                        if column_index in {2, 3, 4}
                        else "center"
                    ),
                    vertical="center",
                    wrap_text=True,
                )

                if column_index in {7, 8}:
                    cell.number_format = '#,##0"원"'

            sheet.row_dimensions[excel_row].height = 32

        widths = {
            1: 7,
            2: 18,
            3: 34,
            4: 28,
            5: 11,
            6: 11,
            7: 14,
            8: 16,
        }
        for column_index, width in widths.items():
            sheet.column_dimensions[
                get_column_letter(column_index)
            ].width = width

        if aggregated_items:
            sheet.auto_filter.ref = (
                f"A4:H{4 + len(aggregated_items)}"
            )

    @staticmethod
    def _write_detail_sheet(
        *,
        sheet,
        supplier_name: str,
        purchase_round: str,
        items: list[dict[str, Any]],
        thin_border: Border,
        header_fill: PatternFill,
    ) -> None:
        sheet.merge_cells("A1:K1")
        sheet["A1"] = "비선상회 주문별 배송목록"
        sheet["A1"].font = Font(size=18, bold=True)
        sheet["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        sheet.row_dimensions[1].height = 32

        sheet["A2"] = "공급처"
        sheet["B2"] = supplier_name
        sheet["D2"] = "발주차수"
        sheet["E2"] = purchase_round
        sheet["G2"] = "생성일시"
        sheet["H2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        sheet["A3"] = "총 주문상품"
        sheet["B3"] = len(items)
        sheet["D3"] = "총 수량"
        sheet["E3"] = sum(
            int(
                item.get(
                    "purchase_quantity",
                    item.get("quantity") or 0,
                )
            )
            for item in items
        )

        headers = [
            "번호",
            "주문번호",
            "공급처 상품명",
            "옵션",
            "수량",
            "수령인",
            "연락처",
            "우편번호",
            "주소",
            "배송메시지",
            "주문일시",
        ]

        for column_index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=4, column=column_index, value=header)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for row_index, item in enumerate(items, start=1):
            excel_row = 4 + row_index
            product_name = (
                item.get("supplier_product_name")
                or item.get("product_name")
                or item.get("platform_product_name")
                or ""
            )

            values = [
                row_index,
                item.get("order_number"),
                product_name,
                item.get("option_name"),
                item.get(
                    "purchase_quantity",
                    item.get("quantity"),
                ),
                item.get("receiver_name"),
                item.get("receiver_phone"),
                item.get("postal_code"),
                full_address(
                    item.get("address"),
                    item.get("detail_address"),
                ),
                item.get("delivery_message"),
                item.get("ordered_at"),
            ]

            for column_index, value in enumerate(values, start=1):
                cell = sheet.cell(
                    row=excel_row,
                    column=column_index,
                    value=value or "",
                )
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal=(
                        "center"
                        if column_index in {
                            1, 2, 5, 6, 7, 8, 11
                        }
                        else "left"
                    ),
                    vertical="center",
                    wrap_text=True,
                )

            sheet.row_dimensions[excel_row].height = 38

        widths = {
            1: 7,
            2: 18,
            3: 28,
            4: 28,
            5: 8,
            6: 11,
            7: 16,
            8: 10,
            9: 48,
            10: 38,
            11: 19,
        }
        for column_index, width in widths.items():
            sheet.column_dimensions[
                get_column_letter(column_index)
            ].width = width

        if items:
            sheet.auto_filter.ref = (
                f"A4:K{4 + len(items)}"
            )
