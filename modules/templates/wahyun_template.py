from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .base import PurchaseTemplate
from .helpers import full_address, get_unique_output_path, safe_filename


class WahyunFarmPurchaseTemplate(PurchaseTemplate):
    """와현농원 전용 발주서 템플릿"""

    template_key = "wahyun"
    display_name = "와현농원 발주서"

    HEADERS = [
        "받으시는 분",
        "받으시는 분 전화",
        "받는분우편번호",
        "받는분총주소",
        "송장출력수량",
        "품목명",
        "운임Type ( C / D)",
        "지불조건 (선불)",
        "특기사항",
        "메모1",
        "보내는분",
        "보내는분 연락처",
    ]

    def create_excel(
        self,
        *,
        output_directory: Path,
        supplier_name: str,
        purchase_round: str,
        items: list[dict[str, Any]],
    ) -> Path:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"

        for col_index, header in enumerate(self.HEADERS, start=1):
            cell = sheet.cell(row=1, column=col_index, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_index, item in enumerate(items, start=2):
            # A receiver_name
            sheet.cell(row=row_index, column=1, value=item.get("receiver_name") or "")
            # B receiver_phone
            sheet.cell(row=row_index, column=2, value=item.get("receiver_phone") or "")
            # C postal_code
            sheet.cell(row=row_index, column=3, value=item.get("postal_code") or "")
            # D full address
            sheet.cell(row=row_index, column=4, value=full_address(item.get("address"), item.get("detail_address")) or "")

            # E: 송장출력수량 - preserve original intent; default to 1 if unknown
            e_value = item.get("label_print_count")
            if e_value is None:
                # do not inject purchase_quantity by default; fall back to 1
                e_value = 1
            sheet.cell(row=row_index, column=5, value=e_value)

            # F: 품목명 supplier_product_name or fallback product_name
            product = item.get("supplier_product_name") or item.get("product_name") or ""
            sheet.cell(row=row_index, column=6, value=product)

            # G: 운임Type mapping based on purchase_quantity
            pq = int(item.get("purchase_quantity") or item.get("quantity") or 0)
            if pq == 1:
                g_val = "c"
            elif pq == 2:
                g_val = "d"
            elif pq >= 3:
                raise ValueError("purchase_quantity >= 3 is not supported for 운임Type")
            else:
                g_val = ""
            sheet.cell(row=row_index, column=7, value=g_val)

            # H: 지불조건 (선불) fixed
            sheet.cell(row=row_index, column=8, value="선불")

            # I: 특기사항 blank
            sheet.cell(row=row_index, column=9, value="")

            # J: 메모1 = delivery_message
            sheet.cell(row=row_index, column=10, value=item.get("delivery_message") or "")

            # K: fixed sender name
            sheet.cell(row=row_index, column=11, value="비선상회")

            # L: fixed sender phone
            sheet.cell(row=row_index, column=12, value="010-2431-0204")

            sheet.row_dimensions[row_index].height = 20

        output_directory.mkdir(parents=True, exist_ok=True)
        created_date = datetime.now().strftime("%Y%m%d")
        file_name = f"{created_date}_더유_와현농원.xlsx"
        file_path = get_unique_output_path(output_directory / file_name)
        workbook.save(file_path)
        return file_path
