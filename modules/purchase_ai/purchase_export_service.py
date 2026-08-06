"""공급처별 발주계획 실무 출력 서비스."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .purchase_ai_repository import PurchaseAIRepository


class PurchaseExportService:
    def __init__(self, db_path: str | Path | None = None) -> None:
        self.repository = PurchaseAIRepository(db_path)

    def get_export_data(self, plan_id: int) -> dict[str, Any]:
        with self.repository.connect() as connection:
            plan = connection.execute(
                "SELECT * FROM purchase_plan WHERE id=?", (plan_id,)
            ).fetchone()
            if not plan:
                raise ValueError("발주계획을 찾을 수 없습니다.")
            summary = connection.execute("""
                SELECT product_code, product_name, COALESCE(option_name,'') AS option_name,
                       purchase_price, SUM(quantity) AS quantity,
                       SUM(estimated_amount) AS estimated_amount,
                       COUNT(DISTINCT order_id) AS order_count
                FROM purchase_plan_item
                WHERE plan_id=? AND excluded=0
                GROUP BY product_code, product_name, option_name, purchase_price
                ORDER BY product_name, option_name
            """, (plan_id,)).fetchall()
            details = connection.execute("""
                SELECT pi.order_number, pi.product_name, COALESCE(pi.option_name,'') AS option_name,
                       pi.quantity, pi.purchase_price, pi.estimated_amount,
                       COALESCE(o.receiver_name,'') AS receiver_name,
                       COALESCE(o.receiver_phone,'') AS receiver_phone,
                       COALESCE(o.postal_code,'') AS postal_code,
                       TRIM(COALESCE(o.address,'') || ' ' || COALESCE(o.detail_address,'')) AS address,
                       COALESCE(o.delivery_message,'') AS delivery_message
                FROM purchase_plan_item pi
                JOIN orders o ON o.id=pi.order_id
                WHERE pi.plan_id=? AND pi.excluded=0
                ORDER BY pi.order_number, pi.product_name, pi.option_name
            """, (plan_id,)).fetchall()
        return {
            "plan": dict(plan),
            "summary": [dict(row) for row in summary],
            "details": [dict(row) for row in details],
        }

    def build_message_text(self, plan_id: int, *, include_delivery: bool = False) -> str:
        data = self.get_export_data(plan_id)
        plan = data["plan"]
        lines = [
            f"[{plan['supplier_name']} 발주 요청]",
            f"작성일: {datetime.now():%Y-%m-%d %H:%M}",
            "",
        ]
        for index, row in enumerate(data["summary"], 1):
            option = f" / {row['option_name']}" if row["option_name"] else ""
            lines.append(f"{index}. {row['product_name']}{option} - {int(row['quantity'] or 0):,}")
        lines += ["", f"총 수량: {int(plan['total_quantity'] or 0):,}"]
        if int(plan.get("estimated_amount", 0) or 0) > 0:
            lines.append(f"예상 매입금액: {int(plan['estimated_amount'] or 0):,}원")
        if include_delivery:
            lines += ["", "[배송 상세]"]
            for index, row in enumerate(data["details"], 1):
                option = f" / {row['option_name']}" if row["option_name"] else ""
                lines += [
                    f"{index}) {row['order_number']} · {row['product_name']}{option} · {int(row['quantity'] or 0):,}",
                    f"   {row['receiver_name']} / {row['receiver_phone']}",
                    f"   ({row['postal_code']}) {row['address']}",
                ]
                if row["delivery_message"]:
                    lines.append(f"   배송메모: {row['delivery_message']}")
        return "\n".join(lines).strip()

    def export_excel(self, plan_id: int, output_path: str | Path) -> Path:
        data = self.get_export_data(plan_id)
        plan = data["plan"]
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "발주요약"
        title_fill = PatternFill("solid", fgColor="D9EAD3")
        header_fill = PatternFill("solid", fgColor="EAF2F8")
        ws["A1"] = f"{plan['supplier_name']} 발주서"
        ws["A1"].font = Font(size=16, bold=True)
        ws.merge_cells("A1:G1")
        ws["A2"] = "작성일"
        ws["B2"] = datetime.now().strftime("%Y-%m-%d %H:%M")
        ws["D2"] = "상태"
        ws["E2"] = plan["status"]
        headers = ["번호", "상품코드", "상품명", "옵션", "수량", "매입단가", "예상금액"]
        for col, value in enumerate(headers, 1):
            cell = ws.cell(4, col, value)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for row_no, row in enumerate(data["summary"], 5):
            values = [row_no - 4, row["product_code"], row["product_name"], row["option_name"],
                      int(row["quantity"] or 0), int(row["purchase_price"] or 0), int(row["estimated_amount"] or 0)]
            for col, value in enumerate(values, 1):
                ws.cell(row_no, col, value)
        total_row = 5 + len(data["summary"])
        ws.cell(total_row, 4, "합계").font = Font(bold=True)
        ws.cell(total_row, 5, int(plan["total_quantity"] or 0)).font = Font(bold=True)
        ws.cell(total_row, 7, int(plan["estimated_amount"] or 0)).font = Font(bold=True)
        for col in range(4, 8):
            ws.cell(total_row, col).fill = title_fill
        for row in range(5, total_row + 1):
            ws.cell(row, 6).number_format = '#,##0"원"'
            ws.cell(row, 7).number_format = '#,##0"원"'

        detail = wb.create_sheet("배송상세")
        detail_headers = ["주문번호", "상품명", "옵션", "수량", "수취인", "연락처", "우편번호", "주소", "배송메모"]
        for col, value in enumerate(detail_headers, 1):
            cell = detail.cell(1, col, value)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for row_no, row in enumerate(data["details"], 2):
            values = [row["order_number"], row["product_name"], row["option_name"], int(row["quantity"] or 0),
                      row["receiver_name"], row["receiver_phone"], row["postal_code"], row["address"], row["delivery_message"]]
            for col, value in enumerate(values, 1):
                detail.cell(row_no, col, value)

        widths = [8, 16, 28, 18, 10, 14, 14]
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        detail_widths = [18, 28, 18, 8, 12, 16, 10, 48, 30]
        for idx, width in enumerate(detail_widths, 1):
            detail.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A5"
        detail.freeze_panes = "A2"
        wb.save(path)
        return path
